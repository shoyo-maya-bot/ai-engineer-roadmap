"""p08_ingest.py のテスト（PDF・Excel・Markdown を Markdown にそろえる. 文書はテストの中で作る）."""

import datetime as dt
from pathlib import Path

from openpyxl import Workbook

from roadmap_examples.p08_ingest import EMPTY_PAGE, cell_text, ingest_dir, main, pdf_to_markdown, xlsx_to_markdown


def make_pdf(pages: list[list[str]]) -> bytes:
    """文字だけの小さな PDF を作る（英数字だけ. テスト用）. 空の一覧のページは、文字の無いページになる."""
    objects: list[bytes] = []
    kids = []
    for i, lines in enumerate(pages):
        content_no = 4 + 2 * i
        page_no = content_no + 1
        ops = "".join(f"BT /F1 12 Tf 72 {720 - 20 * n} Td ({line}) Tj ET\n" for n, line in enumerate(lines))
        stream = ops.encode("ascii")
        objects.append(b"<< /Length %d >>\nstream\n%sendstream" % (len(stream), stream))
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> "
            b"/Contents %d 0 R >>" % content_no
        )
        kids.append(f"{page_no} 0 R")
    head = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(head + objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(offsets) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(offsets) + 1, xref)
    return bytes(out)


def test_pdf_to_markdown_adds_page_headings(tmp_path: Path) -> None:
    path = tmp_path / "error-codes.pdf"
    path.write_bytes(make_pdf([["Error codes", "E-101 invalid input"], []]))
    markdown = pdf_to_markdown(path)
    assert markdown.startswith("# error-codes\n")
    assert "## ページ 1" in markdown
    assert "E-101 invalid input" in markdown
    assert f"## ページ 2\n\n{EMPTY_PAGE}" in markdown  # 文字の無いページ（スキャンした紙など）


def test_xlsx_to_markdown_makes_a_table_per_sheet(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "エラーコード"
    ws.append(["コード", "意味", None])
    ws.append(["E-101", "入力の形が|正しくない", None])
    ws.append([None, None, None])  # 空の行は飛ばす
    ws.append(["E-201", "トークンが\n切れている", None])
    other = wb.create_sheet("更新履歴")
    other.append(["日付", "版"])
    other.append([dt.datetime(2026, 4, 1), 2.0])
    wb.create_sheet("空")
    path = tmp_path / "codes.xlsx"
    wb.save(path)
    markdown = xlsx_to_markdown(path)
    assert "## エラーコード\n\n| コード | 意味 |\n|---|---|\n| E-101 | 入力の形が\\|正しくない |" in markdown
    assert "| E-201 | トークンが 切れている |" in markdown
    assert "| 2026-04-01 | 2 |" in markdown
    assert "## 空\n\n（空のシート）" in markdown


def test_cell_text() -> None:
    assert cell_text(None) == ""
    assert cell_text(3.0) == "3"
    assert cell_text(3.5) == "3.5"
    assert cell_text(dt.datetime(2026, 4, 1, 9, 30)) == "2026-04-01T09:30:00"


def test_ingest_dir_converts_and_reports_problems(tmp_path: Path) -> None:
    src, dst = tmp_path / "src", tmp_path / "dst"
    src.mkdir()
    (src / "overview.md").write_text("# 概要\n\n本文\n", encoding="utf-8")
    (src / "page.pdf").write_bytes(make_pdf([["hello"]]))
    (src / "broken.pdf").write_bytes(b"not a pdf")
    (src / "memo.docx").write_bytes(b"")
    written, problems = ingest_dir(src, dst)
    assert sorted(p.name for p in written) == ["overview.md", "page.md"]
    assert (dst / "overview.md").read_text(encoding="utf-8") == "# 概要\n\n本文\n"
    assert any(p.startswith("broken.pdf: 読めませんでした") for p in problems)
    assert any(p.startswith("memo.docx: 取り込めない形式") for p in problems)
    assert main([str(src), str(tmp_path / "dst2")]) == 1


def test_xlsx_merged_cells_and_unsaved_formulas_lose_values(tmp_path: Path) -> None:
    """取り込みで情報が落ちる例: 結合したセルは左上だけに値が残り、Excel で保存していない数式の値は空になる."""
    wb = Workbook()
    ws = wb.active
    ws.title = "結合"
    ws.append(["項目", "説明", "備考"])
    ws.append(["貸出期間", "14 日", "延長は 1 回"])
    ws.merge_cells("A2:B2")
    formula = wb.create_sheet("数式")
    formula.append(["合計", "=1+1"])
    path = tmp_path / "lost.xlsx"
    wb.save(path)
    markdown = xlsx_to_markdown(path)
    assert "| 貸出期間 |  | 延長は 1 回 |" in markdown  # 「14 日」は結合で消えた
    assert "## 数式\n\n| 合計 |\n|---|" in markdown  # 数式のセルの値は空（右端の空の列は落ちる）
