"""08-1 文書の取り込み: PDF・Excel・Markdown を、検索できる Markdown にそろえる.

- PDF は pypdf でページごとに文字を取り出す（https://pypdf.readthedocs.io/en/stable/user/extract-text.html ）.
  画像だけのページ（スキャンした紙）からは文字を取り出せない. その場合は OCR が要る
- Excel は openpyxl で、シートごとに Markdown の表にする（https://openpyxl.readthedocs.io/en/stable/tutorial.html ）.
  数式のセルは、最後に Excel で保存したときの値を読む（data_only=True）
- Markdown はそのまま写す

使い方（examples/ で）:
  uv run python -m roadmap_examples.p08_ingest <取り込む文書のフォルダ> <Markdown を書くフォルダ>

終了コード: 0 成功 / 1 読めない文書があった / 2 引数の誤り
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

EMPTY_PAGE = "（このページからは文字を取り出せませんでした。画像だけのページなら OCR が要ります）"


def pdf_to_markdown(path: Path) -> str:
    """ページごとに「## ページ N」の見出しを付ける（出典にページを示せるように）."""
    reader = PdfReader(path)
    parts = [f"# {path.stem}"]
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        parts.append(f"## ページ {number}\n\n{text or EMPTY_PAGE}")
    return "\n\n".join(parts) + "\n"


def cell_text(value: Any) -> str:
    """セルの値を、Markdown の表の 1 マスにする."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, dt.datetime) and value.time() == dt.time(0, 0):
        value = value.date()
    if isinstance(value, dt.date | dt.datetime):
        return value.isoformat()
    text = str(value).replace("\r\n", "\n").replace("\n", " ")
    return text.replace("|", "\\|").strip()


def rows_to_table(rows: list[list[str]]) -> str:
    """1 行目を見出しにした Markdown の表. 右端の空の列は落とす."""
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in rows), default=0)
    if width == 0:
        return ""
    rows = [(r + [""] * width)[:width] for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def xlsx_to_markdown(path: Path) -> str:
    """シートごとに「## シート名」の見出しと表を書く. 空の行は飛ばす."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        parts = [f"# {path.stem}"]
        for sheet in workbook.worksheets:
            rows = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            table = rows_to_table(rows)
            parts.append(f"## {sheet.title}\n\n{table or '（空のシート）'}")
    finally:
        workbook.close()  # 読み取り専用のブックは、明示して閉じる（公式ドキュメント）
    return "\n\n".join(parts) + "\n"


CONVERTERS = {
    ".pdf": pdf_to_markdown,
    ".xlsx": xlsx_to_markdown,
    ".md": lambda path: path.read_text(encoding="utf-8-sig"),
}


def ingest_dir(src: Path, dst: Path) -> tuple[list[Path], list[str]]:
    """src の文書を Markdown にして dst に書く. 返り値は（書いたファイル, 読めなかった文書と理由）."""
    dst.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    problems: list[str] = []
    for path in sorted(p for p in src.iterdir() if p.is_file()):
        convert = CONVERTERS.get(path.suffix.lower())
        if convert is None:
            problems.append(f"{path.name}: 取り込めない形式です（{', '.join(CONVERTERS)} だけ）")
            continue
        try:
            markdown = convert(path)
        except Exception as e:  # noqa: BLE001 — 壊れたファイルが 1 つあっても、ほかの文書は取り込む
            problems.append(f"{path.name}: 読めませんでした（{type(e).__name__}: {e}）")
            continue
        out = dst / f"{path.stem}.md"
        out.write_text(markdown, encoding="utf-8")
        written.append(out)
    return written, problems


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="p08_ingest", description="PDF・Excel・Markdown を Markdown にそろえる")
    parser.add_argument("src", type=Path, help="取り込む文書のフォルダ")
    parser.add_argument("dst", type=Path, help="Markdown を書くフォルダ")
    args = parser.parse_args(argv)
    if not args.src.is_dir():
        parser.error(f"フォルダがありません: {args.src}")
    written, problems = ingest_dir(args.src, args.dst)
    for problem in problems:
        print(f"注意: {problem}", file=sys.stderr)
    print(f"{len(written)} 件を書きました", file=sys.stderr)
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
